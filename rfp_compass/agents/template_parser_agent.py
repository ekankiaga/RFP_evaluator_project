"""Parse structured and unstructured evaluation templates for analyst review."""

from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path
from typing import Any

import yaml
from pydantic import ValidationError

from ..models import CriterionDefinition, TemplateReview
from ..tools.document_tool import DocumentTool
from ..llm.key_vault import redact


class TemplateParserAgent:
    """Parse supported files, validate proposals, and compute pack diffs."""
    REQUIRED = ("name", "weight", "max_score")

    def __init__(self, provider: Any | None = None):
        self.provider = provider

    def parse(self, content: bytes, filename: str, domain_pack: dict[str, Any] | None = None) -> TemplateReview:
        extension = Path(filename).suffix.casefold()
        raw_issues: list[str] = []
        try:
            if extension == ".csv":
                rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
            elif extension == ".json":
                value = json.loads(content.decode("utf-8-sig"))
                rows = value.get("criteria", value) if isinstance(value, dict) else value
            elif extension in {".yaml", ".yml"}:
                value = yaml.safe_load(content.decode("utf-8-sig"))
                rows = value.get("criteria", value) if isinstance(value, dict) else value
            elif extension == ".xlsx":
                from openpyxl import load_workbook
                workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                sheet = workbook.active
                matrix = list(sheet.iter_rows(values_only=True))
                headers = [str(x or "").strip() for x in matrix[0]] if matrix else []
                rows = [dict(zip(headers, row)) for row in matrix[1:] if any(x is not None for x in row)]
            elif extension == ".docx":
                from docx import Document
                document = Document(io.BytesIO(content))
                paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
                rows = self._extract_unstructured("\n".join(paragraphs), raw_issues)
                for table in document.tables:
                    matrix = [[cell.text.strip() for cell in row.cells] for row in table.rows]
                    if matrix:
                        rows.extend(dict(zip(matrix[0], row)) for row in matrix[1:])
            elif extension == ".pdf":
                document = DocumentTool().extract(content)
                rows = self._extract_unstructured(document.text, raw_issues)
                raw_issues.extend(document.warnings)
            else:
                raise ValueError("Supported formats are CSV, XLSX, JSON, YAML, DOCX, and PDF.")
        except Exception as exc:
            return TemplateReview(filename=filename, proposed_criteria=[], issues=[redact(f"Could not parse template: {type(exc).__name__}: {exc}")], valid=False)
        if not isinstance(rows, list):
            rows = []
        proposed: list[CriterionDefinition] = []
        for index, row in enumerate(rows, start=1):
            if not isinstance(row, dict):
                raw_issues.append(f"Row {index} is not a criterion object and was skipped.")
                continue
            normalized = self._normalize_row(row)
            try:
                proposed.append(CriterionDefinition.model_validate(normalized))
            except ValidationError as exc:
                raw_issues.append(f"Row {index} is incomplete or invalid: {exc.errors()[0]['msg']}.")
        issues = list(raw_issues)
        self._validate_set(proposed, issues)
        diff = self._diff(proposed, (domain_pack or {}).get("resolved_criteria", []))
        return TemplateReview(filename=filename, proposed_criteria=proposed, issues=issues, diff=diff, valid=not issues and bool(proposed))

    @staticmethod
    def _normalize_row(row: dict[str, Any]) -> dict[str, Any]:
        aliases = {"criterion": "name", "criterion_name": "name", "max": "max_score", "maximum_score": "max_score",
                   "weight_(%)": "weight", "evidence_to_look_for": "what_to_inspect",
                   "mandatory_gate": "is_mandatory_gate", "gate": "is_mandatory_gate", "guidance": "what_to_inspect",
                   "red_flag": "red_flags"}
        values = {}
        for key, value in row.items():
            raw_key = str(key).strip().casefold()
            normalized_key = re.sub(r"[^a-z0-9]+", "_", raw_key).strip("_")
            canonical_key = aliases.get(normalized_key, normalized_key)
            values[canonical_key] = value
        for key in ("weight", "max_score"):
            if key in values and values[key] not in (None, ""):
                values[key] = float(str(values[key]).replace("%", "").strip())
            elif key == "max_score" and key in values:
                values.pop(key)  # An empty maximum uses the standard score range of 0–10.
        for key in ("red_flags", "evidence_hints"):
            value = values.get(key, [])
            if value in (None, ""):
                values[key] = []
            elif isinstance(value, str):
                values[key] = [part.strip() for part in re.split(r"[|;\n]", value) if part.strip()]
        gate = values.get("is_mandatory_gate", False)
        values["is_mandatory_gate"] = gate if isinstance(gate, bool) else str(gate).strip().casefold() in {"true", "yes", "1", "required"}
        return values

    def _extract_unstructured(self, text: str, issues: list[str]) -> list[dict[str, Any]]:
        if not self.provider:
            issues.append("Unstructured template requires an LLM connection; extracted text is not converted automatically.")
            return []
        system = "Extract proposed procurement evaluation criteria. Return JSON only with criteria array and issues array. Preserve only source facts."
        result = self.provider.generate_json(system, text, {"type": "object"})
        payload = json.loads(result.text)
        issues.extend(payload.get("issues", []))
        return payload.get("criteria", [])

    @staticmethod
    def _validate_set(criteria: list[CriterionDefinition], issues: list[str]) -> None:
        names = [c.name.strip().casefold() for c in criteria]
        if len(names) != len(set(names)):
            issues.append("Criterion names must be unique (case-insensitive).")
        if criteria and abs(sum(c.weight for c in criteria) - 100) > 1e-6:
            issues.append(f"Criterion weights total {sum(c.weight for c in criteria):g}%; they must total 100%.")
        for criterion in criteria:
            if len(criterion.name.split()) < 2 or not criterion.description.strip():
                issues.append(f"Criterion '{criterion.name}' may be ambiguous; add a specific name and description.")

    @staticmethod
    def _diff(proposed: list[CriterionDefinition], pack: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
        before = {str(c["name"]).casefold(): c for c in pack}
        after = {c.name.casefold(): c.model_dump() for c in proposed}
        added = [after[name] for name in after.keys() - before.keys()]
        removed = [before[name] for name in before.keys() - after.keys()]
        changed = []
        for key in after.keys() & before.keys():
            old, new = before[key], after[key]
            fields = {field: {"before": old.get(field), "after": new.get(field)} for field in ("description", "weight", "max_score") if old.get(field) != new.get(field)}
            if fields:
                changed.append({"name": new["name"], "fields": fields})
        return {"added": sorted(added, key=lambda x: x["name"].casefold()), "changed": sorted(changed, key=lambda x: x["name"].casefold()), "removed": sorted(removed, key=lambda x: x["name"].casefold())}
