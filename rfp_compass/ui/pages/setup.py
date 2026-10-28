"""Domain, criteria, template review, and template downloads."""

from __future__ import annotations

import io
import json
from pathlib import Path

import streamlit as st
import yaml

from ...agents.template_parser_agent import TemplateParserAgent
from ...agents.domain_classifier import DomainClassifier
from ...audit.audit_log import AuditLog
from ...tools.document_tool import DocumentTool
from ...tools.criteria_tool import CriteriaTool
from ...llm.key_vault import redact


def render() -> None:
    # Apply widget state queued by the previous rerun before the keyed widget is
    # instantiated. Streamlit rejects writes to a widget key later in the run.
    pending_mode = st.session_state.pop("_pending_setup_merge_mode", None)
    if pending_mode in {"Generic + Domain pack", "Generic + Uploaded template", "Template only", "Fully custom"}:
        st.session_state["setup_merge_mode"] = pending_mode
    viewer = st.session_state.get("role") == "Viewer"
    st.title("Setup")
    st.caption("Three steps to a first evaluation: choose a domain, confirm the scoring criteria, then add proposals on the Suppliers page.")
    packs = _packs()
    domain_ids = [pack["domain_id"] for pack in packs]
    st.session_state.setdefault("domain_id", "it_software_saas")
    if st.session_state["domain_id"] not in domain_ids:
        st.session_state["domain_id"] = domain_ids[0]
    current_domain = st.session_state["domain_id"]
    selected_domain = st.selectbox("1 · Choose a procurement domain", domain_ids,
        index=domain_ids.index(current_domain), format_func=lambda domain_id: next(p["name"] for p in packs if p["domain_id"] == domain_id),
        key="setup_domain_choice", disabled=viewer,
        help="Choose the domain that best matches these proposals. You can change it at any time before evaluation.")
    if selected_domain != current_domain:
        st.session_state["domain_id"] = selected_domain
        st.session_state["merge_mode"] = "Generic + Domain pack"
        st.session_state["setup_merge_mode"] = "Generic + Domain pack"
        st.session_state["criteria"] = CriteriaTool().resolve_domain(selected_domain)["resolved_criteria"]
        st.session_state["criteria_domain_id"] = selected_domain
        st.session_state.pop("criteria_set_id", None)
        st.session_state["criteria_version"] = st.session_state.get("criteria_version", 0) + 1
        AuditLog().append(actor=st.session_state.get("actor", "analyst"), role=st.session_state.get("role", "Analyst"),
            event_type="domain_selected", entity_type="domain_pack", entity_id=selected_domain,
            before={"domain_id": current_domain}, after={"domain_id": selected_domain})
    chosen = next(p for p in packs if p["domain_id"] == st.session_state["domain_id"])
    if st.session_state.get("criteria_domain_id") != chosen["domain_id"] or "criteria" not in st.session_state:
        st.session_state["merge_mode"] = "Generic + Domain pack"
        st.session_state["setup_merge_mode"] = "Generic + Domain pack"
        st.session_state["criteria"] = CriteriaTool().resolve_domain(chosen["domain_id"])["resolved_criteria"]
        st.session_state["criteria_domain_id"] = chosen["domain_id"]
    st.info(f"Selected **{chosen['name']}** · {chosen['description']}")

    mode_options = ["Generic + Domain pack", "Generic + Uploaded template", "Template only", "Fully custom"]
    mode_labels = {
        "Generic + Domain pack": "Domain defaults (recommended)",
        "Generic + Uploaded template": "Generic criteria + my uploaded template",
        "Template only": "Use my uploaded template only",
        "Fully custom": "Build a custom scorecard",
    }
    st.session_state.setdefault("merge_mode", "Generic + Domain pack")
    st.subheader("2 · Choose your criteria setup")
    mode = st.selectbox("How should the scorecard be built?", mode_options,
        index=mode_options.index(st.session_state["merge_mode"]), key="setup_merge_mode",
        format_func=lambda value: mode_labels[value], disabled=viewer)
    if mode != st.session_state["merge_mode"]:
        st.session_state["merge_mode"] = mode
        st.session_state["criteria"] = _criteria_for_mode(chosen, mode)
        st.session_state["criteria_version"] = st.session_state.get("criteria_version", 0) + 1
    st.caption({
        "Generic + Domain pack": "Starts with the built-in generic criteria and criteria for the selected procurement domain.",
        "Generic + Uploaded template": "Keeps the generic criteria and adds criteria from your template. Upload and review the template below.",
        "Template only": "Replaces built-in criteria with the criteria in your uploaded template.",
        "Fully custom": "Edit the scorecard rows directly below.",
    }[mode])
    split_key = f"weight_split_{chosen['domain_id']}"
    if mode in {"Generic + Domain pack", "Generic + Uploaded template"}:
        groups = {item.get("group") for item in st.session_state["criteria"] if item.get("is_active", True)}
        if "generic" in groups and groups.intersection({"domain", "template"}):
            st.slider("Weight split · generic criteria vs domain/template criteria", 0, 100,
                int(st.session_state.get(split_key, chosen.get("weight_split", {}).get("generic", 60))),
                key=split_key, on_change=_update_active_split, args=(split_key,), disabled=viewer,
                help="Choose what share of the full 100% belongs to generic criteria. The rest goes to domain or template criteria.")
            split = int(st.session_state[split_key])
            st.caption(f"Generic criteria: **{split}%** · Domain/template criteria: **{100-split}%** · Total: **100%**")
        elif mode == "Generic + Uploaded template":
            st.info("Upload and review a template below to add its criteria. Until then, only the generic criteria are active and their weights total 100%.")

    st.subheader("3 · Review scoring criteria")
    st.caption("Edit individual weights if needed. The total must equal 100% before saving or evaluating.")
    criteria = st.session_state["criteria"]
    editor_key = f"criteria_editor_{chosen['domain_id']}_{st.session_state.get('criteria_version', 0)}"
    edited = st.data_editor(criteria, key=editor_key, use_container_width=True, num_rows="dynamic",
        column_config={"criterion_id": None, "group": None, "red_flags": None, "evidence_hints": None,
                       "sort_order": None, "is_active": None,
                       "weight": st.column_config.NumberColumn("Weight (%)", min_value=0, max_value=100, step=1),
                       "max_score": st.column_config.NumberColumn("Maximum score", min_value=0.1, step=0.5),
                       "what_to_inspect": st.column_config.TextColumn("Evidence to look for"),
                       "is_mandatory_gate": st.column_config.CheckboxColumn("Mandatory gate?")},
        disabled=True if viewer else (["group"] if criteria and "group" in criteria[0] else []))
    st.session_state["criteria"] = edited.to_dict("records") if hasattr(edited, "to_dict") else list(edited)
    total = sum(float(c.get("weight", 0)) for c in st.session_state["criteria"] if c.get("is_active", True))
    a, b, c = st.columns([1, 1, 2])
    a.metric("Weight total", f"{total:.2f}%", delta="Valid" if abs(total - 100) < 1e-6 else "Needs normalization")
    if b.button("Normalize weights to 100%", disabled=viewer or not st.session_state["criteria"]):
        st.session_state["criteria"] = CriteriaTool.normalize_weights(st.session_state["criteria"])
        st.session_state["criteria_version"] = st.session_state.get("criteria_version", 0) + 1
        st.rerun()
    if c.button("Save criteria version (optional)", type="primary", disabled=viewer or abs(total - 100) >= 1e-6):
        try:
            sid = CriteriaTool().snapshot(f"{chosen['name']} criteria", chosen["domain_id"], st.session_state["criteria"], mode)
            AuditLog().append(actor=st.session_state.get("actor", "analyst"), role=st.session_state.get("role", "Analyst"),
                event_type="criteria_set_saved", entity_type="criteria_set", entity_id=sid,
                after={"criteria": st.session_state["criteria"], "merge_mode": mode, "total_weight": total})
            st.session_state["criteria_set_id"] = sid
            st.success(f"Saved criteria set version as `{sid}`.")
        except Exception as exc:
            st.error(str(exc))
    if st.session_state["criteria"]:
        st.caption("Each bar shows the share of the total 100% evaluation weight assigned to that criterion.")
        st.plotly_chart(_weight_chart(st.session_state["criteria"]), use_container_width=True,
                        key=f"setup_weights_{chosen['domain_id']}_{st.session_state.get('criteria_version', 0)}")
    st.divider()
    if mode in {"Generic + Uploaded template", "Template only"}:
        _template_controls(chosen)
    else:
        with st.expander("Optional · use an existing scoring template", expanded=False):
            _template_controls(chosen)
    with st.expander("Optional · AI domain suggestion", expanded=False):
        _domain_suggestion(packs)
    st.success("Next: open **Suppliers**, upload the matching proposal PDFs, then run evaluation in **Evaluate**.")


def _update_active_split(split_key: str) -> None:
    """Apply the slider split and force the editor to render the updated weights."""
    criteria = st.session_state.get("criteria", [])
    _apply_split(criteria, int(st.session_state.get(split_key, 60)))
    st.session_state["criteria"] = criteria
    st.session_state["criteria_version"] = st.session_state.get("criteria_version", 0) + 1


def _template_controls(pack: dict) -> None:
    st.subheader("Use an existing scoring template")
    st.caption("If you are creating a new evaluation, keep the built-in criteria above. Upload a template only when your organization already has its own scorecard.")
    st.info("One row = one criterion. Give each row a weight percentage; all weights must add to 100%. Each criterion is scored from 0 to its maximum score (usually 10).")
    st.markdown("**Simple columns**")
    st.dataframe([
        {"Column": "Criterion name", "What to enter": "A short name, such as Security and privacy."},
        {"Column": "Description", "What to enter": "What a good supplier response should demonstrate."},
        {"Column": "Weight (%)", "What to enter": "Importance of this criterion. All rows together must equal 100."},
        {"Column": "Maximum score", "What to enter": "Usually 10. Scores run from zero up to this number."},
        {"Column": "Evidence to look for", "What to enter": "Documents or details evaluators should find in the proposal."},
        {"Column": "Red flags", "What to enter": "Optional concerns, separated by semicolons."},
        {"Column": "Mandatory gate", "What to enter": "Use true only when failing this item should block award."},
    ], hide_index=True, use_container_width=True)
    uploaded = st.file_uploader("Upload scoring template", type=["csv", "xlsx", "json", "yaml", "yml", "docx", "pdf"], key="template_upload", disabled=st.session_state.get("role") == "Viewer")
    if uploaded:
        import hashlib
        file_hash = hashlib.sha256(uploaded.getvalue()).hexdigest()
        extension = Path(uploaded.name).suffix.casefold()
        unstructured = extension in {".docx", ".pdf"}
        provider = None
        may_parse = not unstructured
        if unstructured:
            connection = st.session_state.get("evaluation_connection")
            connection_id = st.session_state.get("active_connection_id")
            health = st.session_state.get("connection_health", {}).get(connection_id) if connection_id else None
            if not connection or st.session_state.get("demo_mode") or not health or health.status != "passed":
                st.warning("Word and PDF templates need a tested LLM connection assigned for evaluation. CSV, Excel, JSON, and YAML templates do not need an LLM.")
            else:
                st.warning("Only the extracted text from this template will be sent to the selected provider to identify criteria. Supplier proposals are not sent from this page.")
                consent = st.checkbox("I authorize sending this template text to the selected provider", key=f"template_llm_consent_{file_hash}")
                if consent:
                    from ...llm.provider_factory import create_provider
                    provider = create_provider(connection)
                    may_parse = True
        if may_parse and st.session_state.get("template_upload_hash") != file_hash:
            review = TemplateParserAgent(provider).parse(uploaded.getvalue(), uploaded.name, {"resolved_criteria": st.session_state.get("criteria", [])})
            st.session_state["template_review"] = review
            st.session_state["template_upload_hash"] = file_hash
            AuditLog().append(actor=st.session_state.get("actor", "analyst"), role=st.session_state.get("role", "Analyst"),
                event_type="template_uploaded", entity_type="template", entity_id=uploaded.name,
                after={"file_hash": file_hash, "valid": review.valid, "issues": review.issues})
        elif unstructured and not may_parse:
            st.caption("The file is waiting for your consent and a tested connection. No provider request has been made.")
    review = st.session_state.get("template_review")
    if review:
        st.markdown("#### Review the imported criteria")
        diff = review.diff
        st.caption(f"Compared with the current criteria: {len(diff.get('added', []))} added · {len(diff.get('changed', []))} changed · {len(diff.get('removed', []))} removed.")
        if diff.get("added"):
            st.markdown("**New criteria**")
            st.dataframe([{key: item.get(key) for key in ("name", "description", "weight", "max_score")} for item in diff["added"]], hide_index=True, use_container_width=True)
        if diff.get("changed"):
            st.markdown("**Changed criteria**")
            st.dataframe(diff["changed"], hide_index=True, use_container_width=True)
        if diff.get("removed"):
            st.markdown("**Criteria that would be removed**")
            st.dataframe([{key: item.get(key) for key in ("name", "description", "weight")} for item in diff["removed"]], hide_index=True, use_container_width=True)
        table = [item.model_dump() for item in review.proposed_criteria]
        if table:
            review_key = st.session_state.get("template_upload_hash", review.filename)
            edited = st.data_editor(table, key=f"template_review_editor_{review_key}", use_container_width=True, num_rows="dynamic",
                column_config={"criterion_id": None, "is_active": None, "is_mandatory_gate": st.column_config.CheckboxColumn("Mandatory gate?"),
                    "weight": st.column_config.NumberColumn("Weight (%)", min_value=0, max_value=100, step=1),
                    "max_score": st.column_config.NumberColumn("Maximum score", min_value=0.1, step=0.5),
                    "what_to_inspect": st.column_config.TextColumn("Evidence to look for")},
                disabled=st.session_state.get("role") == "Viewer")
            proposed = edited.to_dict("records") if hasattr(edited, "to_dict") else list(edited)
            template_total = sum(float(item.get("weight", 0)) for item in proposed if item.get("is_active", True))
            st.metric("Template weights", f"{template_total:.2f}%", delta="Ready" if abs(template_total - 100) < 1e-6 else "Template rows must total 100%")
            merge_mode = st.session_state.get("merge_mode", "Template only")
            combined = proposed
            duplicate_names: list[str] = []
            if merge_mode == "Generic + Uploaded template":
                split_key = f"weight_split_{pack['domain_id']}"
                initial_split = int(st.session_state.get(split_key, pack.get("weight_split", {}).get("generic", 60)))
                generic_share = st.slider("How should the final 100% be split? · generic vs template", 0, 100,
                    initial_split, key=f"template_review_split_{pack['domain_id']}", disabled=st.session_state.get("role") == "Viewer")
                st.caption(f"Generic criteria: **{generic_share}%** · Uploaded template criteria: **{100-generic_share}%**. Template rows with the same name as a generic criterion update that row instead of creating a duplicate.")
                combined = _combine_generic_template(pack, proposed, generic_share)
                names = [str(item.get("name", "")).casefold().strip() for item in combined if item.get("is_active", True)]
                duplicate_names = sorted({name for name in names if names.count(name) > 1})
                if duplicate_names:
                    st.warning("Some template criteria have the same names as generic criteria. Rename the duplicates in the table above before applying: " + ", ".join(duplicate_names))
                st.caption("Preview of the combined scorecard")
                st.dataframe([{key: item.get(key) for key in ("name", "group", "weight")} for item in combined], hide_index=True, use_container_width=True)
                preview_total = sum(float(item.get("weight", 0)) for item in combined if item.get("is_active", True))
                st.metric("Combined scorecard total", f"{preview_total:.2f}%", delta="Ready" if abs(preview_total - 100) < 1e-6 else "Needs adjustment")
            if review.issues:
                st.markdown("**Please fix these items before accepting:**")
                for issue in review.issues:
                    st.warning(issue)
            ready = abs(template_total - 100) < 1e-6 and abs(sum(float(item.get("weight", 0)) for item in combined if item.get("is_active", True)) - 100) < 1e-6 and not review.issues and not duplicate_names
            if st.button("Apply these criteria", type="primary", disabled=st.session_state.get("role") == "Viewer" or not ready):
                try:
                    from ...models import CriterionDefinition
                    fields = set(CriterionDefinition.model_fields)
                    accepted = []
                    for item in combined:
                        validated = CriterionDefinition.model_validate({key: value for key, value in item.items() if key in fields}).model_dump()
                        if item.get("group"):
                            validated["group"] = item["group"]
                        accepted.append(validated)
                    set_mode = "Generic + Uploaded template" if merge_mode == "Generic + Uploaded template" else "Template only"
                    set_id = CriteriaTool().snapshot(f"{pack['name']} uploaded template", pack["domain_id"], accepted, set_mode)
                    st.session_state["criteria"] = accepted
                    st.session_state["criteria_set_id"] = set_id
                    st.session_state["criteria_version"] = st.session_state.get("criteria_version", 0) + 1
                    if merge_mode == "Generic + Uploaded template":
                        split_key = f"weight_split_{pack['domain_id']}"
                        st.session_state[split_key] = int(st.session_state.get(f"template_review_split_{pack['domain_id']}", 60))
                    AuditLog().append(actor=st.session_state.get("actor", "analyst"), role=st.session_state.get("role", "Analyst"),
                        event_type="template_accepted", entity_type="criteria_set", entity_id=set_id,
                        after={"filename": review.filename, "criteria": accepted, "merge_mode": set_mode})
                    st.session_state["template_review"] = None
                    st.session_state["merge_mode"] = set_mode
                    st.session_state["_pending_setup_merge_mode"] = set_mode
                    st.success("Criteria applied. You can continue to Supplier proposals.")
                    st.rerun()
                except Exception as exc:
                    st.error(redact(exc))
        if st.button("Discard imported template", key="reject_template", disabled=st.session_state.get("role") == "Viewer"):
            st.session_state["template_review"] = None
            st.session_state.pop("template_upload_hash", None)
            st.rerun()
    left, right = st.columns(2)
    left.download_button("Download JSON starter", _blank_json(), "rfp_compass_template.json", "application/json")
    right.download_button("Download Excel starter", _blank_xlsx(), "rfp_compass_template.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

def _packs() -> list[dict]:
    root = Path(__file__).resolve().parents[3] / "domain_packs"
    return [yaml.safe_load(path.read_text(encoding="utf-8")) for path in sorted(root.glob("*.yaml"))]


def _criteria_for_mode(pack: dict, mode: str) -> list[dict]:
    tool = CriteriaTool()
    if mode == "Generic + Domain pack":
        return tool.resolve_domain(pack["domain_id"], mode)["resolved_criteria"]
    generic = [dict(c, group="generic") for c in pack.get("generic_criteria", [])]
    review = st.session_state.get("template_review")
    template = [dict(c.model_dump(), group="template") for c in review.proposed_criteria] if review else []
    if mode == "Generic + Uploaded template":
        if template:
            split_key = f"weight_split_{pack['domain_id']}"
            split = int(st.session_state.get(split_key, pack.get("weight_split", {}).get("generic", 60)))
            return _combine_generic_template(pack, template, split)
        return tool.normalize_weights(generic) if generic else []
    if mode == "Template only":
        return tool.normalize_weights(template) if template else []
    return st.session_state.get("criteria", tool.resolve_domain("custom_generic")["resolved_criteria"])


def _combine_generic_template(pack: dict, template: list[dict], generic_percent: int) -> list[dict]:
    """Merge matching generic rows once and add only new template criteria."""
    combined = [dict(item, group="generic") for item in pack.get("generic_criteria", [])]
    generic_index = {str(item.get("name", "")).casefold().strip(): index for index, item in enumerate(combined)}
    for item in template:
        key = str(item.get("name", "")).casefold().strip()
        if key in generic_index:
            index = generic_index[key]
            # The uploaded row can customize a built-in generic criterion, but
            # it remains in the generic share and appears only once.
            combined[index] = {**combined[index], **item, "group": "generic"}
        else:
            combined.append({**item, "group": "template"})
    _apply_split(combined, generic_percent)
    return combined


def _apply_split(criteria: list[dict], generic_percent: int) -> None:
    generic = [c for c in criteria if c.get("group") == "generic"]
    specific = [c for c in criteria if c.get("group") != "generic"]
    for group, target in ((generic, float(generic_percent)), (specific, float(100 - generic_percent))):
        total = sum(float(c.get("weight", 0)) for c in group)
        if not group or total <= 0:
            continue
        consumed = 0.0
        for index, item in enumerate(group):
            item["weight"] = round(target - consumed, 6) if index == len(group) - 1 else round(float(item["weight"]) * target / total, 6)
            consumed += float(item["weight"])


def _domain_suggestion(packs: list[dict]) -> None:
    st.subheader("Suggest a domain · optional AI assist")
    files = st.session_state.get("supplier_inputs", [])
    if not files:
        st.caption("Upload proposal PDFs on the Suppliers page to enable domain suggestions.")
        return
    if st.button("Suggest domain from proposal text", disabled=st.session_state.get("role") == "Viewer"):
        connection = st.session_state.get("evaluation_connection")
        if not connection or st.session_state.get("demo_mode"):
            st.warning("Add and test an LLM connection first; Demo mode does not classify domains.")
            return
        try:
            from ...llm.provider_factory import create_provider
            extracted = [DocumentTool().extract(item["pdf"]).text for item in files[:3]]
            suggestion = DomainClassifier(create_provider(connection)).suggest("\n".join(extracted), packs)
            st.session_state["domain_suggestion"] = suggestion
        except Exception as exc:
            st.error(f"Domain suggestion failed: {redact(exc)}")
    suggestion = st.session_state.get("domain_suggestion")
    if suggestion:
        st.info(f"Suggested: **{suggestion.domain_id}** · confidence {suggestion.confidence:.0%} · {suggestion.rationale}")
        if st.button("Confirm suggested domain", key="confirm_domain_suggestion", disabled=st.session_state.get("role") == "Viewer"):
            selected = next(p for p in packs if p["domain_id"] == suggestion.domain_id)
            st.session_state["domain_id"] = suggestion.domain_id
            st.session_state["criteria"] = CriteriaTool().resolve_domain(suggestion.domain_id)["resolved_criteria"]
            st.session_state.pop("criteria_set_id", None)
            AuditLog().append(actor=st.session_state.get("actor", "analyst"), role=st.session_state.get("role", "Analyst"),
                event_type="domain_suggestion_confirmed", entity_type="domain_pack", entity_id=suggestion.domain_id,
                after={"confidence": suggestion.confidence, "rationale": suggestion.rationale})
            st.session_state["domain_suggestion"] = None
            st.rerun()


def _weight_chart(criteria: list[dict]):
    import plotly.express as px
    active = [c for c in criteria if c.get("is_active", True)]
    figure = px.bar(active, x="weight", y="name", orientation="h", text="weight",
        labels={"weight": "Share of total weight (%)", "name": "Criterion"},
        color_discrete_sequence=["#39836f"])
    figure.update_traces(texttemplate="%{x:.1f}%", textposition="outside", cliponaxis=False,
        hovertemplate="%{y}<br>%{x:.1f}% of total weight<extra></extra>")
    figure.update_layout(height=max(280, len(active) * 38), margin={"l": 12, "r": 48, "t": 12, "b": 35},
        xaxis={"range": [0, max(105, max((float(c.get("weight", 0)) for c in active), default=0) + 8)], "title": "Share of total weight (%)"},
        yaxis={"title": "", "automargin": True, "autorange": "reversed"}, showlegend=False,
        paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)")
    figure.update_xaxes(gridcolor="#e6ece9", zeroline=False)
    figure.update_yaxes(showgrid=False)
    return figure


def _blank_json() -> bytes:
    """Return a clear, ready-to-edit starter with example weights totaling 100."""
    sample = {"criteria": [
        {"name": "Security and privacy", "description": "Assess how the supplier protects customer and business data.", "weight": 40,
         "max_score": 10, "what_to_inspect": "Encryption, access controls, incident response, and privacy commitments.",
         "red_flags": ["No encryption details", "No incident response process"], "is_mandatory_gate": False},
        {"name": "Implementation plan", "description": "Assess whether the supplier can deliver a practical and well-managed rollout.", "weight": 35,
         "max_score": 10, "what_to_inspect": "Timeline, named responsibilities, dependencies, testing, and migration steps.",
         "red_flags": ["No delivery timeline"], "is_mandatory_gate": False},
        {"name": "Support and service", "description": "Assess the quality and availability of ongoing support.", "weight": 25,
         "max_score": 10, "what_to_inspect": "Support hours, response targets, escalation path, and service reporting.",
         "red_flags": [], "is_mandatory_gate": False},
    ]}
    return json.dumps(sample, indent=2).encode("utf-8")


def _blank_xlsx() -> bytes:
    """Build an Excel starter workbook with criterion examples and a plain-language guide."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Criteria"
    headers = ["Criterion name", "Description", "Weight (%)", "Maximum score", "Evidence to look for", "Red flags", "Mandatory gate"]
    sheet.append(headers)
    examples = [
        ["Security and privacy", "Assess how the supplier protects customer and business data.", 40, 10,
         "Encryption; access controls; incident response; privacy commitments", "No encryption details; no incident response process", "false"],
        ["Implementation plan", "Assess whether the supplier can deliver a practical, well-managed rollout.", 35, 10,
         "Timeline; responsibilities; dependencies; testing; migration steps", "No delivery timeline", "false"],
        ["Support and service", "Assess the quality and availability of ongoing support.", 25, 10,
         "Support hours; response targets; escalation path; service reporting", "", "false"],
    ]
    for row in examples:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:G{sheet.max_row}"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6B64")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 32
    widths = [26, 54, 14, 17, 58, 42, 17]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[2].height = 48
    sheet.row_dimensions[3].height = 48
    sheet.row_dimensions[4].height = 48

    guide = workbook.create_sheet("Read me")
    guide.append(["RFP Compass scoring template"])
    guide.append(["1. Keep one criterion per row on the Criteria sheet."])
    guide.append(["2. Replace the three examples with your own criteria, or add/remove rows."])
    guide.append(["3. The Weight (%) column must add up to exactly 100%."])
    guide.append(["4. Scores run from zero up to Maximum score; 10 is a common maximum."])
    guide.append(["5. Evidence to look for helps reviewers know what proposal text to check."])
    guide.append(["6. Red flags are optional; separate multiple flags with semicolons."])
    guide.append(["7. Set Mandatory gate to true only when a failure should block award."])
    guide.append(["8. Save as .xlsx and upload it in Setup > Use an existing scoring template."])
    guide.column_dimensions["A"].width = 110
    guide["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    guide["A1"].fill = PatternFill("solid", fgColor="1F6B64")
    for row in range(2, guide.max_row + 1):
        guide.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        guide.row_dimensions[row].height = 26
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
